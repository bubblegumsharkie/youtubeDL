from openpyxl import Workbook, load_workbook
from xlsutils import check_existing_id, find_first_inactive_video, set_video_active, get_workbook, set_video_error
from youtubelistreciver import get_list_of_video_ids
from const import LIST_NAME, LIMIT
from downloader import download
from datetime import datetime
from pytubefix.exceptions import AgeRestrictedError

import sys

def save_videos_to_xls(playlist_id):
    all_videos = get_list_of_video_ids(playlist_id)
    # Создаем новый файл Excel или загружаем существующий
    # добавить скачивалку названия плейлиста и сделать именование таблицы
    wb = get_workbook()
    ws = wb.active

    # Создаем заголовки
    ws['A1'] = 'Video Title'
    ws['B1'] = 'Video ID'
    ws['C1'] = 'Downloaded'
    ws['D1'] = 'Added on'

    # Список для отслеживания уже записанных video_id
    existing_ids = []
    # Проходим по столбцу 'B' (Video ID) и добавляем значения в список existing_ids
    for row in ws.iter_rows(min_row=2, max_col=2, max_row=ws.max_row, values_only=True):
        if row[1] is not None:
            existing_ids.append(row[1])

    # Данные для записи
    for video in all_videos:
        video_title = video['snippet']['title']
        video_id = video['snippet']['resourceId']['videoId']

        # Проверяем, существует ли уже video_id
        if not check_existing_id(video_id, existing_ids):
            # Если video_id не существует, добавляем его в список existing_ids
            existing_ids.append(video_id)

            # Добавляем данные в файл Excel
            
            dt_string = datetime.now().strftime("%d/%m/%Y %H:%M")
            row = (video_title, video_id, False, dt_string)  # По умолчанию Is Active - False
            ws.append(row)

    # Сохраняем файл
    wb.save('videos.xlsx')

    print("Данные успешно записаны в файл videos.xlsx.")

def download_video_from_xls():
    wb = get_workbook()
    ws = wb.active
    count = 0
    limit = LIMIT

    while find_first_inactive_video(ws):
        try:
            video_id = find_first_inactive_video(ws)
            link = "https://www.youtube.com/watch?v=" + video_id
            download(link)
            set_video_active(video_id, wb)
            count = count + 1
            if (count > limit):
                print("limit has been reached, please let your system cooldown")
                sys.exit(1)
        except AgeRestrictedError: 
            print("🔴🔴🔴🔴🔴🔴🔴 Download this video manually " + link + " 🔴🔴🔴🔴🔴🔴🔴")
            set_video_error(video_id, wb)

        
# save_videos_to_xls("")
download_video_from_xls()