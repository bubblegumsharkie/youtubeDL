from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

# Проверка, существует ли уже video_id
def check_existing_id(video_id, existing_ids):
    return video_id in existing_ids

# Находим индекс строки с заданным video_id
def find_video_index(video_id, ws):
    for row in range(2, ws.max_row + 1):
        if ws[f'B{row}'].value == video_id:
            return row
    return None

# Устанавливаем значение Downloaded в True для заданного video_id
def set_video_active(video_id, wb):
    ws = wb.active
    index = find_video_index(video_id, ws)
    if index is not None:
        ws[f'C{index}'].value = True
        green_fill = PatternFill(start_color='00FF00', end_color='00FF00', fill_type='solid')
        # Применяем зеленый цвет к строке
        for row in ws.iter_rows(min_row=index, max_row=index):
            for cell in row:
                cell.fill = green_fill
        wb.save('videos.xlsx')

# Устанавливаем значение в Error для заданного video_id
def set_video_error(video_id, wb):
    ws = wb.active
    index = find_video_index(video_id, ws)
    if index is not None:
        ws[f'C{index}'].value = "ERROR"
        ws[f'E{index}'].value = "https://www.youtube.com/watch?v=" + video_id
        wb.save('videos.xlsx')

    # Находим первый video_id, у которого Downloaded == False
def find_first_inactive_video(ws):
    for row in range(2, ws.max_row + 1):
        if not ws[f'C{row}'].value:
            return ws[f'B{row}'].value
    return None

def get_workbook():
    file_path = 'videos.xlsx'
    try:
        wb = load_workbook(filename=file_path)
    except FileNotFoundError:
        wb = Workbook()
    return wb